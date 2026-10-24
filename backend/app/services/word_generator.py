"""Genera el Word de la ficha a partir de `Ficha`, usando `templates/ficha_base.docx`.

La plantilla aporta membrete, pie de página, márgenes y estilos; aquí solo se
escribe el cuerpo: encabezado, foto y las tablas de cada sección.
"""
from __future__ import annotations

import base64
import io
from datetime import date
from xml.sax.saxutils import escape

from docx import Document
from docx.document import Document as DocumentoWord
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from app.config import TEMPLATE_PATH
from app.schemas.ficha import Familiar, Ficha, Vehiculo
from app.services.validator import normalizar

FUENTE = "Palatino Linotype"
COLOR_TEXTO = "000000"
GRIS = "AEAAAA"  # sombreado de los encabezados de tabla (igual al formato original)
ANCHO_TABLA = 8808  # twips (dxa); 1440 = 1 pulgada
COL_ETIQUETA, COL_VALOR = 1686, 7122
SIN_DATO = "Sin dato"

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]  # fmt: skip


def fecha_larga(d: date) -> str:
    return f"{d.day:02d} de {MESES[d.month - 1]} de {d.year}"


def fecha_corta(d: date) -> str:
    return d.strftime("%d/%m/%Y")


def _capitalizar_palabras(texto: str) -> str:
    return texto.title()


# ── Primitivas de formato ────────────────────────────────────────────────────


def _run(p, texto: str, *, negrita=False, cursiva=False, tam=11, versalitas=False):
    r = p.add_run(texto)
    r.bold, r.italic = negrita, cursiva
    r.font.size = Pt(tam)
    r.font.small_caps = versalitas
    r.font.name = FUENTE
    r.font.color.rgb = RGBColor.from_string(COLOR_TEXTO)
    return r


def _formato(p, centrado=False, derecha=False, despues=0):
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(despues), 1.0
    if centrado:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif derecha:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _crear_numeracion_vinetas(doc: DocumentoWord) -> int:
    """Registra una lista con viñetas en numbering.xml y devuelve su numId."""
    numeracion = doc.part.numbering_part.element
    abstractos = [int(a.get(qn("w:abstractNumId"))) for a in numeracion.findall(qn("w:abstractNum"))]
    nums = [int(n.get(qn("w:numId"))) for n in numeracion.findall(qn("w:num"))]
    abs_id, num_id = max(abstractos, default=0) + 1, max(nums, default=0) + 1

    abstracto = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abs_id}">'
        '<w:multiLevelType w:val="hybridMultilevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
        '<w:lvlText w:val="•"/><w:lvlJc w:val="left"/>'
        '<w:pPr><w:ind w:left="284" w:hanging="284"/></w:pPr>'
        f'<w:rPr><w:rFonts w:ascii="{FUENTE}" w:hAnsi="{FUENTE}"/><w:color w:val="{COLOR_TEXTO}"/></w:rPr></w:lvl>'
        "</w:abstractNum>"
    )
    primer_num = numeracion.find(qn("w:num"))  # los abstractNum deben ir antes que los num
    if primer_num is not None:
        primer_num.addprevious(abstracto)
    else:
        numeracion.append(abstracto)

    num = parse_xml(f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abs_id}"/></w:num>')
    cierre = numeracion.find(qn("w:numIdMacAtCleanup"))  # debe quedar al final si existe
    if cierre is not None:
        cierre.addprevious(num)
    else:
        numeracion.append(num)
    return num_id


def _vineta(p, num_id: int) -> None:
    p._p.get_or_add_pPr().insert(
        0,
        parse_xml(
            f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'
        ),
    )


def _hipervinculo(p, url: str, tam=11) -> None:
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    p._p.append(
        parse_xml(
            f'<w:hyperlink {nsdecls("w")} xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            f'r:id="{rid}" w:history="1"><w:r><w:rPr>'
            f'<w:rFonts w:ascii="{FUENTE}" w:hAnsi="{FUENTE}"/><w:i/><w:color w:val="{COLOR_TEXTO}"/>'
            f'<w:sz w:val="{tam * 2}"/><w:u w:val="single"/></w:rPr>'
            f"<w:t>{escape(url)}</w:t></w:r></w:hyperlink>"
        )
    )


def _sombrear(celda, color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    celda._tc.get_or_add_tcPr().append(shd)


def _tabla(doc: DocumentoWord, columnas: list[int]):
    """Tabla centrada, de ancho fijo y con bordes sencillos."""
    t = doc.add_table(rows=0, cols=len(columnas))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    tbl_w = tblPr.find(qn("w:tblW"))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(columnas)))
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{lado}")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "auto")):
            b.set(qn(f"w:{k}"), v)
        bordes.append(b)
    tblPr.find(qn("w:tblLayout")).addprevious(bordes)  # orden exigido por el esquema OOXML
    for gc, ancho in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), columnas):
        gc.set(qn("w:w"), str(ancho))
    return t


def _fila_titulo(t, titulo: str) -> None:
    fila = t.add_row()
    celda = fila.cells[0].merge(fila.cells[-1]) if len(fila.cells) > 1 else fila.cells[0]
    _sombrear(celda, GRIS)
    p = celda.paragraphs[0]
    _formato(p, centrado=True)
    _run(p, titulo, negrita=True)


def _escribir(celda, valores, num_id: int | None = None) -> None:
    """Escribe texto o lista de textos en una celda; con `num_id` usa viñetas."""
    if isinstance(valores, str) or valores is None:
        valores = [valores]
    valores = [v for v in valores if v] or [SIN_DATO]
    for i, texto in enumerate(valores):
        p = celda.paragraphs[0] if i == 0 else celda.add_paragraph()
        if num_id is not None and valores != [SIN_DATO]:
            _vineta(p, num_id)  # antes de _formato: python-docx ordena bien los demás hijos
        _formato(p)
        _run(p, _capitalizar_palabras(texto))


def _fila(t, etiqueta: str, valores, num_id: int | None = None) -> None:
    fila = t.add_row()
    c_etq, c_val = fila.cells
    c_etq.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = c_etq.paragraphs[0]
    _formato(p, centrado=True)
    _run(p, etiqueta, negrita=True)
    _escribir(c_val, valores, num_id)


def _espacio(doc: DocumentoWord) -> None:
    """Párrafo mínimo entre tablas (sin él, Word fusiona tablas contiguas)."""
    p = doc.add_paragraph()
    _formato(p, despues=0)
    # La marca de párrafo también ocupa altura: se reduce a 6 pt.
    p._p.get_or_add_pPr().append(parse_xml(f'<w:rPr {nsdecls("w")}><w:sz w:val="12"/></w:rPr>'))


def _con_fecha(f: Familiar) -> str:
    if f.fecha_nacimiento:
        return f"{f.nombre} (fecha de nacimiento {fecha_corta(f.fecha_nacimiento)})"
    return f.nombre


def _lugar_nacimiento(lugar: str | None, estado: str) -> str:
    if not lugar:
        return estado
    municipio = lugar.strip().rstrip(" ,.")
    ultimo_componente = municipio.rsplit(",", 1)[-1].strip()
    if ultimo_componente.casefold().rstrip(".") == estado.casefold().rstrip("."):
        return municipio
    return f"{municipio}, {estado}"


def _decodificar_foto(b64: str | None) -> bytes | None:
    if not b64:
        return None
    try:
        return base64.b64decode(b64, validate=True)
    except ValueError:
        return None


# ── Secciones ────────────────────────────────────────────────────────────────


def _encabezado(doc: DocumentoWord, f: Ficha, foto: bytes | None) -> None:
    p = doc.add_paragraph()
    _formato(p, derecha=True)
    _run(p, "Fecha: ", negrita=True)
    _run(p, _capitalizar_palabras(fecha_larga(f.fecha)))

    p = doc.add_paragraph()
    _formato(p, derecha=True)
    nuc = f"NUC {f.nuc}" if f.nuc else "NUC sin dato"
    _run(p, nuc if f.nuc else _capitalizar_palabras(nuc), negrita=True)

    p = doc.add_paragraph()
    _formato(p, centrado=True, despues=6)
    _run(p, "Ficha de Identificación Personal", negrita=True, tam=18, versalitas=True)

    p = doc.add_paragraph()
    _formato(p, centrado=True, despues=8)
    _run(p, _capitalizar_palabras(f.datos_personales.nombre or "SIN NOMBRE"), negrita=True, tam=20)

    if foto:
        p = doc.add_paragraph()
        _formato(p, centrado=True, despues=8)
        p.add_run().add_picture(io.BytesIO(foto), width=Inches(1.7))


def _datos_personales(doc, f: Ficha, num_id: int) -> None:
    dp = f.datos_personales
    t = _tabla(doc, [COL_ETIQUETA, COL_VALOR])
    _fila_titulo(t, "DATOS PERSONALES")
    _fila(t, "Nombre:", dp.nombre)
    _fila(t, "Fecha de Nacimiento:", fecha_larga(dp.fecha_nacimiento) if dp.fecha_nacimiento else None)
    _fila(t, "Lugar de nacimiento:", _lugar_nacimiento(dp.lugar_nacimiento, dp.estado))
    _fila(t, "Domicilio:", dp.domicilios, num_id)
    _fila(t, "Teléfono:", dp.telefonos, num_id)
    if dp.telefono_emergencia:
        _fila(t, "Teléfono de emergencia:", f"{dp.telefono_emergencia} (contacto de emergencia)")
    _fila(t, "Ocupación:", dp.ocupacion)
    _fila(t, "Estado Civil:", dp.estado_civil)
    _fila(t, "Edad:", f"{dp.edad} años." if dp.edad is not None else None)
    _fila(t, "Alias:", f'"{dp.alias}"' if dp.alias else None)


def _datos_familiares(doc, f: Ficha, num_id: int) -> None:
    fam = f.datos_familiares
    t = _tabla(doc, [COL_ETIQUETA, COL_VALOR])
    _fila_titulo(t, "DATOS DE FAMILIARES")
    _fila(t, "Madre:", fam.madre)
    _fila(t, "Padre:", fam.padre)
    _fila(t, "Hermanos:", [_con_fecha(x) for x in fam.hermanos], num_id)
    _fila(t, "Hermanas:", [_con_fecha(x) for x in fam.hermanas], num_id)
    _fila(t, "Pareja:", fam.pareja)
    _fila(t, "Hijas:", [_con_fecha(x) for x in fam.hijas])
    _fila(t, "Hijos:", [_con_fecha(x) for x in fam.hijos])


def _carpetas(doc, f: Ficha) -> None:
    t = _tabla(doc, [COL_ETIQUETA, COL_VALOR])
    _fila_titulo(t, "CARPETAS DE INVESTIGACIÓN Y AVERIGUACIONES PREVIAS")
    if not f.carpetas:
        _fila(t, "NUC:", None)
    for i, c in enumerate(f.carpetas):
        if i:  # separador entre carpetas
            sep = t.add_row()
            _sombrear(sep.cells[0].merge(sep.cells[1]), "D9D9D9")
        _fila(t, "NUC:", c.nuc)
        _fila(t, "FECHA INICIO:", fecha_corta(c.fecha_inicio) if c.fecha_inicio else None)
        _fila(t, "DELITO:", c.delito)
        _fila(t, "CALIDAD:", c.calidad)
        _fila(t, "HECHOS:", c.hechos)
        _fila(t, "UNIDAD:", c.unidad)


def _vehiculos(doc, f: Ficha) -> None:
    anchos = [1100, 1500, 1500, 1300, 1500, 1908]
    t = _tabla(doc, anchos)
    _fila_titulo(t, "DATOS VEHICULARES")
    fila = t.add_row()
    for celda, titulo in zip(fila.cells, ["ALTA", "MARCA", "LÍNEA", "MODELO", "PLACA", "SERIE"]):
        p = celda.paragraphs[0]
        _formato(p, centrado=True)
        _run(p, titulo, negrita=True)
    if not f.vehiculos:
        fila = t.add_row()
        _escribir(fila.cells[0].merge(fila.cells[-1]), SIN_DATO)
    for v in f.vehiculos:
        fila = t.add_row()
        for celda, valor in zip(fila.cells, [v.alta, v.marca, v.linea, v.modelo, v.placa, v.serie]):
            _escribir(celda, valor or "")


def _lista_simple(doc, titulo: str, items: list[str], num_id: int) -> None:
    t = _tabla(doc, [ANCHO_TABLA])
    _fila_titulo(t, titulo)
    _escribir(t.add_row().cells[0], items, num_id)


def _fuentes_abiertas(doc, f: Ficha, num_id: int) -> None:
    t = _tabla(doc, [ANCHO_TABLA])
    _fila_titulo(t, "FUENTES ABIERTAS")
    celda = t.add_row().cells[0]
    if not f.fuentes_abiertas:
        _escribir(celda, SIN_DATO)
        return
    primero = True
    for fa in f.fuentes_abiertas:
        p = celda.paragraphs[0] if primero else celda.add_paragraph()
        primero = False
        _vineta(p, num_id)
        _formato(p)
        _run(p, _capitalizar_palabras(fa.descripcion))
        if fa.fecha_consulta:
            p = celda.add_paragraph()
            _formato(p)
            _run(p, _capitalizar_palabras(f"Fecha y hora de consulta: {fa.fecha_consulta}"), cursiva=True)
        if fa.link:
            p = celda.add_paragraph()
            _formato(p, despues=4)
            _run(p, "Link: ", cursiva=True)
            _hipervinculo(p, fa.link)


def _fotografias(doc, f: Ficha, foto: bytes | None, num_id: int) -> None:
    if not foto and not f.fotografias and not f.nota_fotografia:
        return
    t = _tabla(doc, [3600, ANCHO_TABLA - 3600])
    _fila_titulo(t, "FOTOGRAFÍAS")
    c_foto, c_info = t.add_row().cells
    p = c_foto.paragraphs[0]
    _formato(p, centrado=True)
    if foto:
        p.add_run().add_picture(io.BytesIO(foto), width=Inches(2.0))
    primero = True
    if f.nota_fotografia:
        p = c_info.paragraphs[0]
        _formato(p)
        _run(p, "Nota: ", negrita=True)
        _run(p, _capitalizar_palabras(f.nota_fotografia))
        primero = False
    for ft in f.fotografias:
        lineas = []  # (etiqueta, valor, valor_en_negrita)
        if ft.fuente:
            lineas.append(("", ft.fuente, True))
        for etiqueta, valor in (
            ("Fecha: ", fecha_corta(ft.fecha) if ft.fecha else None),
            ("Nombre: ", ft.nombre),
            ("Número de causa: ", ft.numero_causa),
        ):
            if valor:
                lineas.append((etiqueta, valor, False))
        for i, (etiqueta, valor, negrita) in enumerate(lineas):
            p = c_info.paragraphs[0] if primero else c_info.add_paragraph()
            if i == 0:
                _vineta(p, num_id)  # la viñeta marca el inicio de cada registro
            primero = False
            _formato(p)
            if etiqueta:
                _run(p, etiqueta, negrita=True)
            _run(p, _capitalizar_palabras(valor), negrita=negrita)


# ── Punto de entrada ─────────────────────────────────────────────────────────


def generar_docx(ficha: Ficha) -> bytes:
    """Devuelve el contenido del .docx de la ficha."""
    ficha = normalizar(ficha)
    doc = Document(str(TEMPLATE_PATH))
    num_id = _crear_numeracion_vinetas(doc)
    foto = _decodificar_foto(ficha.foto_base64)

    _encabezado(doc, ficha, foto)
    _datos_personales(doc, ficha, num_id)
    _espacio(doc)
    _datos_familiares(doc, ficha, num_id)
    _espacio(doc)
    _carpetas(doc, ficha)
    _espacio(doc)
    _vehiculos(doc, ficha)
    _espacio(doc)
    _lista_simple(doc, "REGISTRO PÚBLICO DE LA PROPIEDAD", ficha.registro_publico, num_id)
    _espacio(doc)
    _fuentes_abiertas(doc, ficha, num_id)
    _espacio(doc)
    _fotografias(doc, ficha, foto, num_id)
    if ficha.observaciones:
        _espacio(doc)
        t = _tabla(doc, [ANCHO_TABLA])
        _fila_titulo(t, "OBSERVACIONES")
        _escribir(t.add_row().cells[0], ficha.observaciones)

    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()
