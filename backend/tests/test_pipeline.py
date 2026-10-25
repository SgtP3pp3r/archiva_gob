import io
from datetime import date
from zipfile import ZipFile
from xml.etree import ElementTree

import pytest
from docx import Document
from fastapi.testclient import TestClient

from main import app
from app.schemas.ficha import Ficha
from app.services.data_extractor import extraer_ficha, parse_fecha
from app.services.pdf_reader import ContenidoPdf, PdfError, leer_pdf
from app.services.validator import calcular_edad, normalizar, validar
from app.services.word_generator import generar_docx
from tests.conftest import crear_pdf

client = TestClient(app)


def test_api_root():
    response = client.get("/")

    assert response.status_code == 200
    assert response.json() == {
        "message": "API de Fichas de Identificación",
        "docs": "/docs",
    }


# ── pdf_reader ───────────────────────────────────────────────────────────────
def test_lee_texto_e_imagenes(pdf_bytes):
    c = leer_pdf(pdf_bytes)
    assert "DATOS PERSONALES" in c.texto and c.paginas == 2
    # Quedan solo la foto y la captura. Se descartan: fondo de página, membrete, pie,
    # sombra plana y la marca de agua repetida en todas las páginas.
    assert sorted((i.ancho, i.alto) for i in c.imagenes) == [(160, 200), (400, 150)]


def test_rechaza_no_pdf():
    with pytest.raises(PdfError):
        leer_pdf(b"hola, esto no es un pdf")


def test_rechaza_pdf_sin_texto():
    with pytest.raises(PdfError, match="OCR"):
        leer_pdf(crear_pdf(lineas=[], con_foto=True))


# ── data_extractor ───────────────────────────────────────────────────────────
@pytest.fixture
def ficha(pdf_bytes) -> Ficha:
    return extraer_ficha(leer_pdf(pdf_bytes))


def test_datos_personales(ficha):
    dp = ficha.datos_personales
    assert ficha.nuc == "19-2019-12345"
    assert dp.nombre == "Juan Carlos Pérez Ramírez"
    assert dp.fecha_nacimiento == date(1993, 9, 11)
    assert dp.estado == "Chihuahua"
    assert dp.edad == 33
    assert dp.alias == "El Ejemplo"
    assert dp.telefonos == ["6561234567"]
    # el domicilio cortado en dos renglones se reconstruye en un solo elemento
    assert len(dp.domicilios) == 2
    assert dp.domicilios[0].endswith("Ciudad Ejemplo, Chihuahua.")


def test_familiares(ficha):
    fam = ficha.datos_familiares
    assert fam.madre == "María López Soto."
    assert [(h.nombre, h.fecha_nacimiento) for h in fam.hermanos] == [
        ("Luis Pérez Ramírez", date(1986, 5, 21)),
        ("Pedro Pérez Ramírez", date(1985, 4, 7)),
    ]
    assert fam.hermanas == [] and fam.hijas == []  # "Sin dato"
    assert fam.hijos[0].nombre == "Mateo Pérez Torres"
    assert fam.pareja == "Ana Torres Gil (Posible pareja)"


def test_extrae_reporte_piweb_y_genera_word():
    texto = (
        "FUENTE: IDENTIDAD - PADRON N, NOMBRE: ANA MARIA, PATERNO: LOPEZ, "
        "MATERNO: RUIZ, FECHA_NACIMIENTO: 10/02/1980, CALLE: JUAREZ, "
        "NUMERO_EXTERIOR: 15, NUM_INT: 4, COLONIA: CENTRO, CP: 31000, "
        "MUNICIPIO: CHIHUAHUA, ESTADO6: SONORA, TELEFONO: 6141234567, "
        "TELEFONO_EMERGENCIA: 6147654321\n"
        "FUENTE: OTRO REGISTRO, NOMBRE: ANA MARIA, PATERNO: LOPEZ, "
        "MATERNO: RUIZ, FECHA_NACIMIENTO: 10/02/1980, TELEFONO: 6149999999\n"
        "FUENTE: REGISTRO CIVIL - NACIMIENTO, FECHA: 10/02/1980, NOMBRE: ANA MARIA, "
        "PATERNO: LOPEZ, MATERNO: RUIZ, NOMBRE_PADRE18: JOSE, PATERNO_PADRE18: "
        "LOPEZ, MATERNO_PADRE18: GARCIA, NOMBRE_MADRE18: ELENA, PATERNO_MADRE18: "
        "RUIZ, MATERNO_MADRE18: SOTO\n"
        "FUENTE: REGISTRO CIVIL - NACIMIENTO, FECHA: 21/05/2005, NOMBRE: LUIS, "
        "PATERNO: LOPEZ, MATERNO: PEREZ, NOMBRE_PADRE18: ANA MARIA, "
        "PATERNO_PADRE18: LOPEZ, MATERNO_PADRE18: RUIZ, NOMBRE_MADRE18: CARLA, "
        "PATERNO_MADRE18: PEREZ, MATERNO_MADRE18: SOTO\n"
        "FUENTE: REGISTRO CIVIL - MATRIMONIO, NOMBRE: ANA MARIA, PATERNO: LOPEZ, "
        "MATERNO: RUIZ, NOMBRE_ESPOSO: CARLOS, PATERNO_ESPOSO: PEREZ, "
        "MATERNO_ESPOSO: SOTO\n"
        "FUENTE: SH - PADRON VEHICULAR, NOMBRE: ANA MARIA, PATERNO: LOPEZ, "
        "MATERNO: RUIZ, FECHA_NACIMIENTO: 10/02/1980, FECHA: 01/03/2020, "
        "PLACA: ABC123, MARCA: NISSAN, LINEA: VERSA, MODELO: 2019\n"
    )
    ficha = normalizar(extraer_ficha(ContenidoPdf(texto=texto, paginas=1)))

    assert ficha.datos_personales.nombre == "ANA MARIA LOPEZ RUIZ"
    assert ficha.datos_personales.fecha_nacimiento == date(1980, 2, 10)
    assert ficha.datos_personales.estado == "SONORA"
    assert ficha.datos_personales.domicilios == [
        "Calle JUAREZ, #4, Colonia CENTRO, C.P. 31000, CHIHUAHUA, SONORA"
    ]
    assert ficha.datos_personales.telefonos == ["6141234567"]
    assert ficha.datos_personales.telefono_emergencia == "6147654321"
    assert ficha.datos_familiares.padre == "JOSE LOPEZ GARCIA"
    assert ficha.datos_familiares.madre == "ELENA RUIZ SOTO"
    assert ficha.datos_familiares.pareja == "CARLOS PEREZ SOTO"
    assert [h.nombre for h in ficha.datos_familiares.hijos] == ["LUIS LOPEZ PEREZ"]
    assert ficha.vehiculos[0].placa == "ABC123"
    assert validar(ficha).valido
    word = Document(io.BytesIO(generar_docx(ficha)))
    texto_word = "\n".join(p.text for p in word.paragraphs)
    celdas_word = "\n".join(c.text for t in word.tables for row in t.rows for c in row.cells)
    assert "Ana Maria Lopez Ruiz" in texto_word
    assert "Calle Juarez, #4, Colonia Centro, C.P. 31000, Chihuahua, Sonora" in celdas_word
    assert "Chihuahua, Sonora" in celdas_word


def test_direccion_piweb_usa_chihuahua_si_no_hay_estado():
    texto = (
        "FUENTE: IDENTIDAD - PADRON N, NOMBRE: ANA, PATERNO: LOPEZ, "
        "MATERNO: RUIZ, FECHA_NACIMIENTO: 10/02/1980, CALLE: JUAREZ, "
        "NUMERO_EXTERIOR: 15, MUNICIPIO: CHIHUAHUA\n"
    )
    ficha = extraer_ficha(ContenidoPdf(texto=texto, paginas=1))

    assert ficha.datos_personales.domicilios == ["Calle JUAREZ, #15, CHIHUAHUA"]


def test_direccion_piweb_separa_colonia_y_cp_incrustados_en_municipio():
    texto = (
        "FUENTE: IDENTIDAD - PADRON N, NOMBRE: ANA, PATERNO: LOPEZ, "
        "MATERNO: RUIZ, FECHA_NACIMIENTO: 10/02/1980, CALLE: JUAREZ, "
        "NUM_INT: 4, MUNICIPIO: CHIHUAHUA, COLONIA CENTRO, C.P. 31000, ESTADO6: CHIHUAHUA\n"
    )
    ficha = extraer_ficha(ContenidoPdf(texto=texto, paginas=1))

    assert ficha.datos_personales.domicilios == [
        "Calle JUAREZ, #4, Colonia CENTRO, C.P. 31000, CHIHUAHUA"
    ]


def test_direccion_piweb_elimina_componentes_repetidos():
    texto = (
        "FUENTE: IDENTIDAD - PADRON N, NOMBRE: ANA, PATERNO: LOPEZ, "
        "MATERNO: RUIZ, CALLE: JUAREZ, COLONIA: CENTRO, MUNICIPIO: CHIHUAHUA, "
        "ESTADO6: CHIHUAHUA\n"
    )
    ficha = extraer_ficha(ContenidoPdf(texto=texto, paginas=1))

    assert ficha.datos_personales.domicilios == [
        "Calle JUAREZ, Colonia CENTRO, CHIHUAHUA"
    ]


def test_extrae_telefono_de_emergencia_en_ficha_tradicional():
    texto = (
        "DATOS PERSONALES\nNombre: Ana López\n"
        "Teléfono de emergencia: 6141234567\n"
    )
    ficha = extraer_ficha(ContenidoPdf(texto=texto, paginas=1))

    assert ficha.datos_personales.telefono_emergencia == "6141234567"


def test_carpetas_y_secciones(ficha):
    assert len(ficha.carpetas) == 1
    c = ficha.carpetas[0]
    assert (c.nuc, c.fecha_inicio, c.calidad) == ("19-2019-12345", date(2019, 10, 13), "Imputado")
    assert c.hechos == "Narración de hechos ficticia que ocupa varias líneas dentro de la carpeta."
    assert ficha.vehiculos == [] and ficha.registro_publico == []
    fa = ficha.fuentes_abiertas[0]
    assert fa.link == "https://www.facebook.com/ejemplo.perez"
    assert fa.fecha_consulta.startswith("09/10/2026 13:27")
    assert fa.descripcion == "Posible página de Facebook de Juan Carlos Pérez Ramírez."
    ft = ficha.fotografias[0]
    assert (ft.fuente, ft.fecha, ft.numero_causa) == ("Registro de Ejemplo", date(2023, 4, 27), "25/2019.")
    assert ficha.observaciones == "Solicitud por medio de oficio FGE-000/2026."


def test_foto_es_el_retrato(ficha):
    import base64, struct
    png = base64.b64decode(ficha.foto_base64)
    ancho, alto = struct.unpack(">II", png[16:24])  # cabecera IHDR del PNG
    assert (ancho, alto) == (160, 200)  # el retrato: no la captura apaisada ni ninguna decoración


def test_parse_fecha():
    assert parse_fecha("09 de Octubre de 2026") == date(2026, 10, 9)
    assert parse_fecha("31/02/2020") is None
    assert parse_fecha("sin fecha") is None


# ── validator ────────────────────────────────────────────────────────────────
def test_edad_se_calcula_si_falta():
    f = Ficha(fecha=date(2026, 10, 9))
    f.datos_personales.fecha_nacimiento = date(1993, 9, 11)
    assert normalizar(f).datos_personales.edad == 33
    assert calcular_edad(date(1993, 10, 10), date(2026, 10, 9)) == 32


def test_validador_detecta_problemas(ficha):
    assert validar(ficha).valido
    ficha.nuc = "19-2019-12354"  # dígitos transpuestos respecto a la carpeta
    ficha.datos_personales.telefonos = ["123"]
    ficha.datos_personales.nombre = None
    r = validar(ficha)
    campos = {i.campo: i.nivel for i in r.incidencias}
    assert not r.valido and campos["datos_personales.nombre"] == "error"
    assert campos["nuc"] == "advertencia"
    assert campos["datos_personales.telefonos"] == "advertencia"


# ── word_generator ───────────────────────────────────────────────────────────
def test_normalizar_filtra_domicilios_y_vehiculos_que_recibe_editar():
    from app.schemas.ficha import Vehiculo

    ficha = Ficha()
    ficha.datos_personales.domicilios = [
        "Calle Juárez, #4, Colonia Centro, C.P. 31000, Chihuahua, Chihuahua",
        "calle JUAREZ, #4, colonia centro, C.P. 31000, CHIHUAHUA, CHIHUAHUA",
        "Calle Hidalgo, #2, Colonia Centro, C.P. 31000, Chihuahua, Chihuahua",
    ]
    ficha.vehiculos = [
        Vehiculo(serie="ABC-123"),
        Vehiculo(serie="abc 123"),
        Vehiculo(serie="XYZ-987"),
    ]

    filtrada = normalizar(ficha)

    assert filtrada.datos_personales.domicilios == [
        "Calle Juárez, #4, Colonia Centro, C.P. 31000, Chihuahua, Chihuahua",
        "Calle Hidalgo, #2, Colonia Centro, C.P. 31000, Chihuahua, Chihuahua",
    ]
    assert [vehiculo.serie for vehiculo in filtrada.vehiculos] == ["ABC-123", "XYZ-987"]
    assert len(ficha.datos_personales.domicilios) == 3
    assert len(ficha.vehiculos) == 3


def test_normalizar_elimina_domicilios_con_valores_sin_nombre_o_numero():
    ficha = Ficha()
    ficha.datos_personales.domicilios = [
        "Calle Sin nombre, #15, Colonia Centro, C.P. 31000, Chihuahua, Chihuahua",
        "Calle Juárez, #S/N, Colonia Centro, C.P. 31000, Chihuahua, Chihuahua",
        "Calle Juárez, #Sin número, Colonia Centro, C.P. 31000, Chihuahua, Chihuahua",
        "Calle Juárez, #4, Colonia Centro, C.P. 31000, Chihuahua, Chihuahua",
    ]

    filtrada = normalizar(ficha)

    assert filtrada.datos_personales.domicilios == [
        "Calle Juárez, #4, Colonia Centro, C.P. 31000, Chihuahua, Chihuahua"
    ]
    assert len(ficha.datos_personales.domicilios) == 4


def test_docx_contiene_los_datos(ficha):
    contenido = generar_docx(ficha)
    doc = Document(io.BytesIO(contenido))
    texto = "\n".join(p.text for p in doc.paragraphs)
    celdas = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Juan Carlos Pérez Ramírez" in texto
    assert "NUC 19-2019-12345" in texto
    for esperado in ("11 De Septiembre De 1993", "6561234567", "Luis Pérez Ramírez (Fecha De Nacimiento 21/05/1986)",
                     "Sin Dato", "FUENTES ABIERTAS", "Imputado", '"El Ejemplo"', "Registro De Ejemplo"):
        assert esperado in celdas
    assert len(doc.inline_shapes) == 2  # foto superior + tabla de fotografías
    with ZipFile(io.BytesIO(contenido)) as archivo:
        xml = ElementTree.fromstring(archivo.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        texto_runs = [
            run for run in xml.findall(".//w:r", ns)
            if run.find(".//w:t", ns) is not None
        ]
        assert texto_runs
        assert all(
            run.find("w:rPr/w:color", ns) is not None
            and run.find("w:rPr/w:color", ns).get(f"{{{ns['w']}}}val") == "000000"
            for run in texto_runs
        )


def test_word_filtra_domicilios_y_vehiculos_duplicados_y_anexa_nota(ficha):
    from app.schemas.ficha import Vehiculo

    ficha.datos_personales.domicilios = [
        "Calle Juárez, #4, Colonia Centro, C.P. 31000, Chihuahua, Chihuahua",
        "calle JUÁREZ, #4, colonia centro, C.P. 31000, CHIHUAHUA, CHIHUAHUA",
        "Calle Hidalgo, #2, Colonia Centro, C.P. 31000, Chihuahua, Chihuahua",
    ]
    ficha.datos_personales.telefono_emergencia = "6141234567"
    ficha.vehiculos = [
        Vehiculo(marca="Nissan", serie="ABC-123"),
        Vehiculo(marca="Nissan", serie="abc 123"),
        Vehiculo(marca="Ford", serie="XYZ-987"),
    ]
    ficha.nota_fotografia = "Fotografía obtenida del registro consultado."

    doc = Document(io.BytesIO(generar_docx(ficha)))
    personales = next(t for t in doc.tables if "DATOS PERSONALES" in t.rows[0].cells[0].text)
    fila_domicilios = next(row for row in personales.rows if row.cells[0].text == "Domicilio:")
    assert len(fila_domicilios.cells[1].text.splitlines()) == 2
    assert "Hidalgo" in fila_domicilios.cells[1].text
    fila_emergencia = next(row for row in personales.rows if "emergencia" in row.cells[0].text.casefold())
    assert "6141234567 (Contacto De Emergencia)" == fila_emergencia.cells[1].text

    vehiculos = next(t for t in doc.tables if "DATOS VEHICULARES" in t.rows[0].cells[0].text)
    series = [row.cells[5].text for row in vehiculos.rows[2:]]
    assert len(series) == 2
    assert series[0] == "Abc-123" and series[1] == "Xyz-987"

    fotografias = next(t for t in doc.tables if "FOTOGRAFÍAS" in t.rows[0].cells[0].text)
    assert "Nota: Fotografía Obtenida Del Registro Consultado." in fotografias.rows[1].cells[1].text


def test_plantilla_no_contiene_datos_de_personas():
    from app.config import TEMPLATE_PATH
    doc = Document(str(TEMPLATE_PATH))
    assert not doc.paragraphs or all(not p.text.strip() for p in doc.paragraphs)
    assert not doc.tables


# ── API ──────────────────────────────────────────────────────────────────────
def test_api_flujo_completo(pdf_bytes):
    r = client.post("/api/documents/extract", files={"archivo": ("f.pdf", pdf_bytes, "application/pdf")})
    assert r.status_code == 200
    cuerpo = r.json()
    assert cuerpo["validacion"]["valido"]
    assert cuerpo["ficha"]["datos_personales"]["domicilios"]
    cuerpo["ficha"]["datos_personales"]["ocupacion"] = "Editado por el usuario"
    r = client.post("/api/documents/generate", json=cuerpo["ficha"])
    assert r.status_code == 200
    assert "Ficha_Juan_Carlos_Perez_Ramirez.docx" in r.headers["content-disposition"]
    doc = Document(io.BytesIO(r.content))
    assert "Editado Por El Usuario" in "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    datos_personales = next(t for t in doc.tables if "DATOS PERSONALES" in t.rows[0].cells[0].text)
    lugar_nacimiento = next(row.cells[1].text for row in datos_personales.rows if row.cells[0].text == "Lugar de nacimiento:")
    assert lugar_nacimiento == "Ciudad Ejemplo, Chihuahua"


def test_api_extract_muestra_solo_domicilios_sin_marcadores_invalidos():
    pdf = crear_pdf(lineas=[
        "DATOS PERSONALES",
        "Nombre: Ana López",
        "Domicilio:",
        "• Calle Sin nombre, #15, Colonia Centro, C.P. 31000, Chihuahua, Chihuahua",
        "• Calle Juárez, #S/N, Colonia Centro, C.P. 31000, Chihuahua, Chihuahua",
        "• Calle Juárez, #4, Colonia Centro, C.P. 31000, Chihuahua, Chihuahua",
    ], con_foto=False)
    response = client.post(
        "/api/documents/extract",
        files={"archivo": ("domicilios.pdf", pdf, "application/pdf")},
    )

    assert response.status_code == 200
    assert response.json()["ficha"]["datos_personales"]["domicilios"] == [
        "Calle Juárez, #4, Colonia Centro, C.P. 31000, Chihuahua, Chihuahua"
    ]


def test_api_atajo_pdf_a_word(pdf_bytes):
    r = client.post("/api/documents/ficha", files={"archivo": ("f.pdf", pdf_bytes, "application/pdf")})
    assert r.status_code == 200 and r.headers["x-errores"] == "0"


def test_api_errores():
    r = client.post("/api/documents/extract", files={"archivo": ("x.pdf", b"no soy pdf", "application/pdf")})
    assert r.status_code == 422
    r = client.post("/api/documents/generate", json={})  # sin nombre → error de validación
    assert r.status_code == 422


def test_nombre_de_archivo_sin_punto_doble(pdf_bytes):
    ficha = extraer_ficha(leer_pdf(pdf_bytes))
    ficha.datos_personales.nombre = "María Pérez López."  # los PDF suelen traer punto final
    r = client.post("/api/documents/generate", json=ficha.model_dump(mode="json"))
    assert ".." not in r.headers["content-disposition"]
    assert "filename*=UTF-8''Ficha_Mar%C3%ADa_P%C3%A9rez_L%C3%B3pez.docx" in r.headers["content-disposition"]
